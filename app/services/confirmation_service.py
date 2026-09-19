from __future__ import annotations

import logging
import uuid
from datetime import datetime, date
from io import BytesIO
from typing import Optional

import openpyxl

from app.dtos.requests.confirmation_request_dto import (
    CreateConfirmationDTO,
    UpdateConfirmationDTO,
)
from app.dtos.responses.confirmation_response_dto import (
    ConfirmationListResponse,
    ConfirmationResponse,
)
from app.enums.excel_headers import ExcelHeader
from app.mappers.confirmation_mapper import (
    confirmation_to_response,
    confirmations_to_list_response,
)
from app.models.confirmation import Confirmation
from app.models.order import Order
from app.repositories.confirmation_repository import ConfirmationRepository
from app.repositories.order_repository import OrderRepository
from app.utils.order_dates import as_date, as_display
from app.services import order_service
from app.configuration.app_config import AppConfig
from app.services.email_service import _extract_provider_number, send_delivery_email
from app.services.pdf_service import download_from_s3, upload_file_to_s3

logger = logging.getLogger(__name__)

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Column name used in crossdocking Excel files for delivery date
CROSSDOCKING_DELIVERY_DATE_COL = "FECHA_ENTREGA"


def create_confirmation(
    organization_id: str, dto: CreateConfirmationDTO
) -> ConfirmationResponse:
    with ConfirmationRepository() as conf_repo, OrderRepository() as order_repo:
        existing = conf_repo.find_by_company_and_number(
            organization_id, dto.confirmation_number
        )
        if existing:
            raise ValueError(
                f"Confirmation '{dto.confirmation_number}' already exists for organization {organization_id}"
            )

        orders = order_repo.find_by_company_and_documents(
            organization_id, dto.document_numbers
        )
        found_numbers = {o.document_number for o in orders}
        missing = set(dto.document_numbers) - found_numbers
        if missing:
            raise LookupError(
                f"Orders not found: {', '.join(sorted(missing))}"
            )

        # Validate and extract delivery place + date from orders
        store_id = _validate_deliver_to(orders, [])
        earliest_date = _validate_order_dates(orders, [])

        confirmation = Confirmation(
            company_id=organization_id,
            confirmation_number=dto.confirmation_number,
            delivery_date=earliest_date or None,
            deliver_to_store_id=store_id,
        )
        confirmation = conf_repo.save(confirmation)

        _link_orders_to_confirmation(orders, confirmation, order_repo)

        confirmation = conf_repo.save(confirmation)
        return confirmation_to_response(confirmation)


def update_confirmation(
    organization_id: str, confirmation_number: str, dto: UpdateConfirmationDTO
) -> ConfirmationResponse:
    with ConfirmationRepository() as conf_repo, OrderRepository() as order_repo:
        confirmation = conf_repo.find_by_company_and_number(
            organization_id, confirmation_number
        )
        if not confirmation:
            raise LookupError(
                f"Confirmation '{confirmation_number}' not found for organization {organization_id}"
            )

        orders = order_repo.find_by_company_and_documents(
            organization_id, dto.document_numbers
        )
        found_numbers = {o.document_number for o in orders}
        missing = set(dto.document_numbers) - found_numbers
        if missing:
            raise LookupError(
                f"Orders not found: {', '.join(sorted(missing))}"
            )

        _link_orders_to_confirmation(orders, confirmation, order_repo)

        confirmation = conf_repo.save(confirmation)
        return confirmation_to_response(confirmation)


def get_confirmation(
    organization_id: str, confirmation_number: str
) -> ConfirmationResponse:
    with ConfirmationRepository() as conf_repo:
        confirmation = conf_repo.find_by_company_and_number(
            organization_id, confirmation_number
        )
        if not confirmation:
            raise LookupError(
                f"Confirmation '{confirmation_number}' not found for organization {organization_id}"
            )
        return confirmation_to_response(confirmation)


def list_confirmations(
    organization_id: str, page: int = 1, page_size: int = 12
) -> ConfirmationListResponse:
    with ConfirmationRepository() as conf_repo:
        confirmations, total = conf_repo.find_all_by_company(
            organization_id, page=page, page_size=page_size
        )
        return confirmations_to_list_response(confirmations, page, page_size, total)


STATUS_MAP = {1: "pending", 2: "processing", 3: "shipped", 4: "delivered", 5: "cancelled"}


def update_confirmation_status(
    organization_id: str, confirmation_number: str, status_code: int
) -> ConfirmationResponse:
    """Update the status for a confirmation and all its orders.

    Status codes: 1=pending, 2=processing, 3=shipped, 4=delivered, 5=cancelled.
    When status is 3 (shipped), sends a delivery email with NuevoReporte attachments.
    """
    status = STATUS_MAP.get(status_code)
    if not status:
        raise ValueError(f"Invalid status code: {status_code}")

    with ConfirmationRepository() as conf_repo, OrderRepository() as order_repo:
        confirmation = conf_repo.find_by_company_and_number(
            organization_id, confirmation_number
        )
        if not confirmation:
            raise LookupError(
                f"Confirmation '{confirmation_number}' not found for organization {organization_id}"
            )

        confirmation.confirmation_status = status

        active_orders = [o for o in (confirmation.orders or []) if o.status == 1]
        for order in active_orders:
            order.order_status = status
            order_repo.save(order)

        # When marking as shipped (3), send the delivery email with reports
        if status_code == 3 and active_orders:
            _send_confirmation_email(confirmation, active_orders)

        confirmation = conf_repo.save(confirmation)
        return confirmation_to_response(confirmation)


def _send_confirmation_email(confirmation: Confirmation, orders: list[Order]) -> None:
    """Collect NuevoReporte Excel attachments and send delivery email."""
    # The email prints it, so format here — the column is a real date.
    delivery_date = as_display(confirmation.delivery_date)

    # Extract provider number from the organization's internal_code
    provider_number = ""
    for order in orders:
        org = order.organization
        if org and org.internal_code:
            provider_number = _extract_provider_number(org.internal_code)
            break

    # Collect NuevoReporte Excel URLs as attachments
    attachments: list[dict] = []
    for order in orders:
        if order.nuevo_reporte_url:
            last4 = (order.document_number or "")[-4:]
            attachments.append({
                "url": order.nuevo_reporte_url,
                "filename": f"{last4}-RN.xlsx",
            })

    if not attachments:
        logger.warning(
            f"No NuevoReporte attachments found for confirmation {confirmation.confirmation_number}"
        )

    try:
        send_delivery_email(delivery_date, provider_number, attachments)
    except Exception as e:
        logger.error(f"Failed to send delivery email for confirmation {confirmation.confirmation_number}: {e}")


def remove_order_from_confirmation(
    organization_id: str, confirmation_number: str, document_number: str
) -> ConfirmationResponse:
    with ConfirmationRepository() as conf_repo, OrderRepository() as order_repo:
        confirmation = conf_repo.find_by_company_and_number(
            organization_id, confirmation_number
        )
        if not confirmation:
            raise LookupError(
                f"Confirmation '{confirmation_number}' not found for organization {organization_id}"
            )

        order = order_repo.find_by_company_and_document(
            organization_id, document_number
        )
        if not order:
            raise LookupError(
                f"Order '{document_number}' not found for organization {organization_id}"
            )

        if order.confirmation_id != confirmation.confirmation_id:
            raise ValueError(
                f"Order '{document_number}' is not linked to confirmation '{confirmation_number}'"
            )

        order.confirmation_id = None
        order.confirmation_number = None
        order_repo.save(order)

        confirmation = conf_repo.find_by_company_and_number(
            organization_id, confirmation_number
        )
        return confirmation_to_response(confirmation)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _parse_date(value) -> Optional[date]:
    """Coerce an order date, whatever shape it is in.

    Was `strptime(..., "%d/%m/%Y")`, which HARD-FAILED on the ISO dates the POS
    and the storefront write — so linking a manual order to a confirmation raised
    "invalid delivery date format" about a perfectly valid date. It now goes
    through the one helper that knows all the shapes, including the real `date`
    the column returns after migration d3e4f5a6b7c8.
    """
    return as_date(value)


def _validate_deliver_to(
    new_orders: list[Order], existing_orders: list[Order]
) -> Optional[int]:
    """Validate all orders share the same delivery store.

    Returns the deliver_to_store_id from the orders, or None.
    """
    all_orders = list(new_orders) + list(existing_orders)
    store_ids = set()
    result_store_id = None

    for order in all_orders:
        if order.deliver_to_store_id:
            store_ids.add(order.deliver_to_store_id)
            result_store_id = order.deliver_to_store_id

    if len(store_ids) > 1:
        codes = []
        for order in all_orders:
            if order.deliver_to_store and order.deliver_to_store.store_code:
                codes.append(order.deliver_to_store.store_code)
        raise ValueError(
            f"All orders in a confirmation must be delivered to the same place. "
            f"Found different delivery stores: {', '.join(sorted(set(codes)))}"
        )

    return result_store_id


def _validate_order_dates(
    new_orders: list[Order], existing_orders: list[Order]
) -> Optional[date]:
    """Validate date rules and return the earliest delivery date.

    Rules:
      1. No order can have a delivery date in the past.
      2. All orders (new + existing) must be in the same month/year.

    Returns the earliest delivery date across all orders, as a `date`.
    """
    today = date.today()
    all_dates: list[date] = []
    all_orders = list(new_orders) + list(existing_orders)

    for order in all_orders:
        if not order.delivery_date:
            continue
        d = _parse_date(order.delivery_date)
        if d is None:
            raise ValueError(
                f"Order '{order.document_number}' has an unreadable delivery date: "
                f"'{order.delivery_date}'"
            )
        if d < today:
            raise ValueError(
                f"Order '{order.document_number}' has a delivery date in the past ({order.delivery_date})"
            )
        all_dates.append(d)

    if not all_dates:
        return None

    # Check all dates share the same month/year
    ref_month = all_dates[0].month
    ref_year = all_dates[0].year
    for d in all_dates[1:]:
        if d.month != ref_month or d.year != ref_year:
            raise ValueError(
                f"All orders in a confirmation must be in the same month. "
                f"Found dates in {all_dates[0].strftime('%m/%Y')} and {d.strftime('%m/%Y')}"
            )

    # A `date`, not a formatted string: the column it is written to is a date
    # now, and formatting belongs at the point of display.
    return min(all_dates)


def _link_orders_to_confirmation(
    orders: list[Order], confirmation: Confirmation, order_repo: OrderRepository
) -> None:
    # Gather already-linked active orders (excluding the ones being added now)
    new_doc_numbers = {o.document_number for o in orders}
    existing_orders = [
        o for o in (confirmation.orders or [])
        if o.status == 1 and o.document_number not in new_doc_numbers
    ]

    # Validate delivery place and dates
    store_id = _validate_deliver_to(orders, existing_orders)
    if store_id:
        confirmation.deliver_to_store_id = store_id

    earliest_date = _validate_order_dates(orders, existing_orders)
    if earliest_date:
        confirmation.delivery_date = earliest_date

    for order in orders:
        if (
            order.confirmation_id is not None
            and order.confirmation_id != confirmation.confirmation_id
        ):
            raise ValueError(
                f"Order '{order.document_number}' is already linked to a different confirmation"
            )

        order.confirmation_id = confirmation.confirmation_id
        order.confirmation_number = confirmation.confirmation_number
        if confirmation.delivery_date:
            order.delivery_date = confirmation.delivery_date

        if confirmation.delivery_date:
            # The spreadsheet keeps DD/MM/YYYY — it is what the chain reads
            # back — so the date is formatted at this boundary rather than stored
            # that way.
            _update_detalles_excel_date(order, as_display(confirmation.delivery_date))
            _update_crossdocking_excel_date(order, as_display(confirmation.delivery_date))

        order_repo.save(order)
        _reprocess_order_safe(order.company_id, order.document_number)

    # Also update existing orders if the earliest date changed
    if earliest_date and existing_orders:
        for order in existing_orders:
            if as_date(order.delivery_date) != earliest_date:
                order.delivery_date = earliest_date
                _update_detalles_excel_date(order, as_display(earliest_date))
                _update_crossdocking_excel_date(order, as_display(earliest_date))
                order_repo.save(order)
                _reprocess_order_safe(order.company_id, order.document_number)


def _reprocess_order_safe(organization_id: str, document_number: str) -> None:
    try:
        order_service.reprocess_order(organization_id, document_number)
    except Exception as e:
        logger.warning(
            f"Failed to reprocess order {document_number}: {e}"
        )


def _update_detalles_excel_date(order: Order, new_date) -> None:
    """Download the DETALLES Excel from S3, update the delivery date column, re-upload."""
    if not order.excel_url:
        return
    try:
        excel_bytes = download_from_s3(order.excel_url)
        wb = openpyxl.load_workbook(BytesIO(excel_bytes))
        ws = wb.active

        # Find the delivery date column using ExcelHeader aliases
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        date_col_idx = None
        for alias in ExcelHeader.DELIVERY_DATE.aliases:
            if alias in headers:
                date_col_idx = headers.index(alias)
                break

        if date_col_idx is None:
            logger.warning(
                f"Delivery date column not found in DETALLES Excel for order {order.document_number}"
            )
            return

        # Update all data rows (row 2 onwards, 1-indexed in openpyxl)
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=date_col_idx + 1, value=new_date)

        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Re-upload to same S3 key
        _reupload_excel(order.excel_url, buf.getvalue())
        logger.info(f"Updated DETALLES Excel delivery date for order {order.document_number}")
    except Exception as e:
        logger.warning(
            f"Failed to update DETALLES Excel for order {order.document_number}: {e}"
        )


def _update_crossdocking_excel_date(order: Order, new_date) -> None:
    """Download the crossdocking Excel from S3, update the FECHA_ENTREGA column, re-upload."""
    if not order.crossdocking_excel_url:
        return
    try:
        excel_bytes = download_from_s3(order.crossdocking_excel_url)
        wb = openpyxl.load_workbook(BytesIO(excel_bytes))
        ws = wb.active

        headers = [str(cell.value or "").strip() for cell in ws[1]]
        date_col_idx = None
        if CROSSDOCKING_DELIVERY_DATE_COL in headers:
            date_col_idx = headers.index(CROSSDOCKING_DELIVERY_DATE_COL)

        if date_col_idx is None:
            logger.warning(
                f"FECHA_ENTREGA column not found in crossdocking Excel for order {order.document_number}"
            )
            return

        # Update the metadata row (row 2)
        ws.cell(row=2, column=date_col_idx + 1, value=new_date)

        buf = BytesIO()
        wb.save(buf)
        wb.close()

        _reupload_excel(order.crossdocking_excel_url, buf.getvalue())
        logger.info(
            f"Updated crossdocking Excel delivery date for order {order.document_number}"
        )
    except Exception as e:
        logger.warning(
            f"Failed to update crossdocking Excel for order {order.document_number}: {e}"
        )


def _reupload_excel(url: str, file_bytes: bytes) -> None:
    """Re-upload an Excel file to the same S3 location derived from its URL."""
    pdf_domain = (AppConfig.get_key("pdf.domain", "") or "").rstrip("/")
    if pdf_domain and url.startswith(pdf_domain):
        key = url[len(pdf_domain):].lstrip("/")
    else:
        key = url.split(".amazonaws.com/", 1)[-1]

    upload_file_to_s3(file_bytes, key, EXCEL_CONTENT_TYPE)
