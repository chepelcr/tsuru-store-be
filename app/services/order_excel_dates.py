"""Write a delivery date back into an order's stored spreadsheets.

**Why a database update alone is not enough.** `order_service.reprocess_order`
re-parses the Excel files held in S3 and `_update_order_from_parsed` assigns
`order.delivery_date` from what it finds — so a date changed only in the database
is silently reverted the next time anyone reprocesses that order, and Reprocess is
an item in the order's own menu in the POS. Whatever changes the date has to
change the spreadsheet too, or the change is temporary in a way nobody is told
about.

This lived inside `confirmation_service`, reachable only as a side effect of
linking orders to a confirmation. It is here so the order edit endpoint and the
confirmation flow run the same code: two implementations of "rewrite the date in
the sheet" is two chances for one of them to miss the crossdocking file.

The cells stay `DD/MM/YYYY`. The columns are real dates now (migration
`d3e4f5a6b7c8`) but the spreadsheet is read by the chain and by our own importer,
and both expect day-first — so the date is formatted at this boundary rather than
stored that way.
"""

from __future__ import annotations

import logging
from io import BytesIO

import openpyxl

from app.configuration.app_config import AppConfig
from app.enums.excel_headers import ExcelHeader
from app.models.order import Order
from app.services.pdf_service import download_from_s3, upload_file_to_s3
from app.utils.order_dates import DateLike, as_display

logger = logging.getLogger(__name__)

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: The crossdocking sheet's own header for the delivery date.
CROSSDOCKING_DELIVERY_DATE_COL = "FECHA_ENTREGA"


def rewrite_delivery_date(order: Order, new_date: DateLike) -> None:
    """Put `new_date` into both of the order's spreadsheets, if it has them.

    Best-effort by design: a failure is logged and swallowed. The database row is
    the source of truth for everything except a reprocess, and refusing the whole
    edit because an S3 object could not be rewritten would block a correction the
    operator can otherwise make. The log is what makes the divergence findable.
    """
    formatted = as_display(new_date)
    if not formatted:
        return
    _update_detalles_excel_date(order, formatted)
    _update_crossdocking_excel_date(order, formatted)


def _update_detalles_excel_date(order: Order, new_date: str) -> None:
    """Rewrite the delivery-date column in the DETALLES sheet, every data row."""
    if not order.excel_url:
        return
    try:
        excel_bytes = download_from_s3(order.excel_url)
        workbook = openpyxl.load_workbook(BytesIO(excel_bytes))
        sheet = workbook.active

        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        column_index = None
        for alias in ExcelHeader.DELIVERY_DATE.aliases:
            if alias in headers:
                column_index = headers.index(alias)
                break

        if column_index is None:
            logger.warning(
                "Delivery date column not found in DETALLES Excel for order %s",
                order.document_number,
            )
            return

        # Every data row: the importer reads the date from the first one, but a
        # sheet whose rows disagree is a sheet somebody will later read wrongly.
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row=row_index, column=column_index + 1, value=new_date)

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        _reupload_excel(order.excel_url, buffer.getvalue())
        logger.info(
            "Updated DETALLES Excel delivery date for order %s", order.document_number)
    except Exception as exc:  # noqa: BLE001
        logger.warning(
            "Failed to update DETALLES Excel for order %s: %s", order.document_number, exc)


def _update_crossdocking_excel_date(order: Order, new_date: str) -> None:
    """Rewrite FECHA_ENTREGA in the crossdocking sheet's metadata row."""
    if not order.crossdocking_excel_url:
        return
    try:
        excel_bytes = download_from_s3(order.crossdocking_excel_url)
        workbook = openpyxl.load_workbook(BytesIO(excel_bytes))
        sheet = workbook.active

        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if CROSSDOCKING_DELIVERY_DATE_COL not in headers:
            logger.warning(
                "FECHA_ENTREGA column not found in crossdocking Excel for order %s",
                order.document_number,
            )
            return
        column_index = headers.index(CROSSDOCKING_DELIVERY_DATE_COL)

        # Metadata row only — this sheet carries the order's header on row 2 and
        # its line items below, and the date belongs to the header.
        sheet.cell(row=2, column=column_index + 1, value=new_date)

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        _reupload_excel(order.crossdocking_excel_url, buffer.getvalue())
        logger.info(
            "Updated crossdocking Excel delivery date for order %s", order.document_number)
    except Exception as exc:  # noqa: BLE001
        logger.warning(
            "Failed to update crossdocking Excel for order %s: %s",
            order.document_number, exc,
        )


def _reupload_excel(url: str, file_bytes: bytes) -> None:
    """Overwrite the same S3 key the URL points at.

    The key is derived from the stored URL rather than rebuilt from the order, so
    the rewrite lands on the object the order actually references — including for
    rows written before the current key convention.
    """
    pdf_domain = (AppConfig.get_key("pdf.domain", "") or "").rstrip("/")
    if pdf_domain and url.startswith(pdf_domain):
        key = url[len(pdf_domain):].lstrip("/")
    else:
        key = url.split(".amazonaws.com/", 1)[-1]

    upload_file_to_s3(file_bytes, key, EXCEL_CONTENT_TYPE)
